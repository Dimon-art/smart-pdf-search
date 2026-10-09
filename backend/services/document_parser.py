"""
Универсальный парсер документов: PDF, DOCX, TXT, MD.

Возвращает единую структуру для chunker/embedder:
{
    "type": "pdf" | "docx" | "txt" | "md",
    "pages": [
        {
            "page": int,
            "text": str,
            "blocks": [{"text": str, "bbox": list[float] | None}, ...]
        },
        ...
    ],
    "supports_highlight": bool
}

Для PDF — реальные страницы + bbox (подсветка работает).
Для DOCX/TXT/MD — блоки без bbox (подсветка недоступна, замена — да).
"""

from __future__ import annotations

import os
from typing import Any

from services.pdf_parser import (
    extract_pages_with_bbox,
    extract_blocks_with_pages,
    clean_text,
    get_pdf_metadata,
)


# ---------------------------------------------------------------------------
# Публичный API
# ---------------------------------------------------------------------------

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def parse_document(path: str) -> dict[str, Any]:
    """Определяет формат по расширению и вызывает нужный парсер."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in (".txt", ".md"):
        return parse_txt(path, ext.lstrip("."))
    raise ValueError(f"Неподдерживаемый формат: {ext}")


def get_document_metadata(path: str) -> dict[str, Any]:
    """Единая мета для всех форматов (name, size, pages, type)."""
    ext = os.path.splitext(path)[1].lower()
    size = os.path.getsize(path) if os.path.exists(path) else 0
    name = os.path.basename(path)

    pages = 0
    if ext == ".pdf" and os.path.exists(path):
        try:
            meta = get_pdf_metadata(path)
            pages = meta.get("pages", 0)
        except Exception:
            pages = 0
    elif ext == ".docx" and os.path.exists(path):
        try:
            import docx
            doc = docx.Document(path)
            pages = _estimate_docx_pages(doc)
        except Exception:
            pages = 0

    return {
        "name": name,
        "size": size,
        "pages": pages,
        "type": ext.lstrip("."),
    }


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def parse_pdf(path: str) -> dict[str, Any]:
    pages_raw = extract_pages_with_bbox(path)
    blocks = extract_blocks_with_pages(path)

    # Группируем блоки по страницам
    blocks_by_page: dict[int, list[dict]] = {}
    for b in blocks:
        blocks_by_page.setdefault(b["page"], []).append(
            {"text": b["text"], "bbox": b.get("bbox")}
        )

    pages_out = []
    for p in pages_raw:
        page_num = p["page"]
        pages_out.append(
            {
                "page": page_num,
                "text": clean_text(p["text"]),
                "blocks": blocks_by_page.get(page_num, []),
            }
        )

    return {
        "type": "pdf",
        "pages": pages_out,
        "supports_highlight": True,
    }


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def parse_docx(path: str) -> dict[str, Any]:
    """
    DOCX не имеет понятия «страница» на уровне API python-docx.
    Разбиваем параграфы на условные страницы по ~40 параграфов,
    чтобы уложиться в структуру, аналогичную PDF.
    """
    import docx

    doc = docx.Document(path)

    # Собираем непустые параграфы, сохраняя текст таблиц
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    # Разбиваем на «страницы»
    PAGE_SIZE = 40
    pages_out = []
    for i in range(0, len(paragraphs), PAGE_SIZE):
        chunk = paragraphs[i : i + PAGE_SIZE]
        page_idx = i // PAGE_SIZE
        pages_out.append(
            {
                "page": page_idx,
                "text": clean_text("\n".join(chunk)),
                "blocks": [{"text": clean_text(p), "bbox": None} for p in chunk],
            }
        )

    if not pages_out:
        pages_out = [{"page": 0, "text": "", "blocks": []}]

    return {
        "type": "docx",
        "pages": pages_out,
        "supports_highlight": False,
    }


def _estimate_docx_pages(doc) -> int:
    n = sum(1 for p in doc.paragraphs if p.text.strip())
    return max(1, (n + 39) // 40)


# ---------------------------------------------------------------------------
# TXT / MD
# ---------------------------------------------------------------------------

def parse_txt(path: str, kind: str = "txt") -> dict[str, Any]:
    """
    TXT/MD читаем как UTF-8 (fallback: cp1251).
    Разбиваем по пустым строкам на блоки, блоки группируем в «страницы».
    """
    text = _read_text_file(path)
    lines = text.splitlines()

    # Блоки = абзацы (разделены пустой строкой)
    blocks_raw: list[str] = []
    buf: list[str] = []
    for line in lines:
        if line.strip() == "":
            if buf:
                blocks_raw.append("\n".join(buf).strip())
                buf = []
        else:
            buf.append(line)
    if buf:
        blocks_raw.append("\n".join(buf).strip())

    # Для MD — срезаем markdown-разметку в блоках (для поиска мешает)
    if kind == "md":
        blocks_raw = [_strip_md(b) for b in blocks_raw]

    PAGE_SIZE = 40
    pages_out = []
    for i in range(0, len(blocks_raw), PAGE_SIZE):
        chunk = blocks_raw[i : i + PAGE_SIZE]
        pages_out.append(
            {
                "page": i // PAGE_SIZE,
                "text": clean_text("\n\n".join(chunk)),
                "blocks": [{"text": clean_text(b), "bbox": None} for b in chunk],
            }
        )

    if not pages_out:
        pages_out = [{"page": 0, "text": "", "blocks": []}]

    return {
        "type": kind,
        "pages": pages_out,
        "supports_highlight": False,
    }


def _read_text_file(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _strip_md(text: str) -> str:
    """Убирает базовую markdown-разметку (заголовки, жирный, курсив, код)."""
    import re
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)   # заголовки
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)                 # **bold**
    text = re.sub(r"__(.+?)__", r"\1", text)                     # __bold__
    text = re.sub(r"\*(.+?)\*", r"\1", text)                     # *italic*
    text = re.sub(r"_(.+?)_", r"\1", text)                       # _italic_
    text = re.sub(r"`(.+?)`", r"\1", text)                       # `code`
    text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)              # [text](url)
    return text.strip()